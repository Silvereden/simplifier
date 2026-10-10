import io
import re
import uuid
from datetime import datetime
from typing import Any, List, Optional

import docx
import pypdf
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.database import get_db
from app.models.models import (
    Customer,
    DeliveryNote,
    DeliveryStatus,
    OrderStatus,
    PurchaseOrder,
)

router = APIRouter()


# --- Utilitaire d'extraction de texte (PDF / Word) ---

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    text = ""
    filename_lower = filename.lower()
    
    if filename_lower.endswith(".pdf"):
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    elif filename_lower.endswith(".docx") or filename_lower.endswith(".doc"):
        doc = docx.Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                text += " ".join([cell.text for cell in row.cells]) + "\n"
    else:
        text = file_bytes.decode("utf-8", errors="ignore")
        
    return text


# --- Schémas Pydantic ---

class OrderItemSchema(BaseModel):
    product_name: str
    quantity: int
    unit_price: float

class OrderCreateSchema(BaseModel):
    customer_id: str
    po_number: str
    items: List[OrderItemSchema]

class EmailOrderItemSchema(BaseModel):
    product_name: str
    quantity: int
    unit_price: float = 0.0

class OrderFromEmailSchema(BaseModel):
    client_name: str
    client_email: Optional[str] = None
    items: List[EmailOrderItemSchema]
    notes: Optional[str] = None


# --- Routes ---

@router.get("/", status_code=status.HTTP_200_OK)
async def get_purchase_orders(db: AsyncSession = Depends(get_db)):
    """0. Récupérer la liste des Bons de Commande (PO) pour le tableau de bord"""
    try:
        result = await db.execute(
            select(PurchaseOrder).options(selectinload(PurchaseOrder.customer))
        )
        orders = result.scalars().all()

        return [
            {
                "id": str(po.id),
                "po_number": po.po_number,
                "status": str(po.status.value) if hasattr(po.status, "value") else str(po.status),
                "total_amount": float(po.total_amount or 0.0),
                "customer": {
                    "company_name": po.customer.company_name
                } if po.customer else None
            }
            for po in orders
        ]
    except Exception as e:
        print(f"Erreur lors du chargement des commandes : {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Erreur serveur : {str(e)}"
        )


@router.post("/", status_code=status.HTTP_201_CREATED)
async def create_purchase_order(
    order_in: OrderCreateSchema, 
    db: AsyncSession = Depends(get_db)
):
    """1. Enregistrer un nouveau Bon de Commande (PO) manuel"""
    total = sum(item.quantity * item.unit_price for item in order_in.items)
    
    po = PurchaseOrder(
        customer_id=order_in.customer_id,
        po_number=order_in.po_number,
        status=OrderStatus.PENDING,
        total_amount=total
    )
    db.add(po)
    await db.commit()
    await db.refresh(po)
    return {"message": "Bon de commande créé avec succès", "order_id": po.id}


@router.post("/from-email", status_code=status.HTTP_201_CREATED)
async def create_purchase_order_from_email(
    payload: OrderFromEmailSchema,
    db: AsyncSession = Depends(get_db)
):
    """2. Réceptionner et enregistrer un PO automatiquement (Make / Webhook / Mailhook)"""
    
    # A. Recherche du client par email ou par nom
    customer = None
    if payload.client_email:
        result = await db.execute(select(Customer).filter(Customer.email == payload.client_email))
        customer = result.scalars().first()
    
    if not customer and payload.client_name:
        result = await db.execute(
            select(Customer).filter(Customer.company_name.ilike(f"%{payload.client_name}%"))
        )
        customer = result.scalars().first()
        
    # B. Auto-création du client s'il est inconnu avec tous les champs non-nulls
    if not customer:
        clean_code = "".join(e for e in (payload.client_name or "CUST") if e.isalnum()).upper()[:6] or "CUST"
        customer_code = f"C-{clean_code}-{str(uuid.uuid4())[:4].upper()}"

        customer = Customer(
            code=customer_code,
            company_name=payload.client_name or "Client Inconnu",
            email=payload.client_email or f"contact@{clean_code.lower()}.com",
            address="Zone Industrielle",
            city="Tétouan",
            ice="000000000000000"
        )
        db.add(customer)
        await db.commit()
        await db.refresh(customer)

    # C. Génération automatique du numéro de PO
    raw_prefix = getattr(customer, "code", None) or customer.company_name[:3]
    prefix = "".join(e for e in raw_prefix if e.isalnum()).upper()
    timestamp = datetime.now().strftime("%Y%m%d-%H%M")
    po_number = f"PO-{prefix}-{timestamp}"

    # D. Calcul du montant total
    total_amount = sum(item.quantity * item.unit_price for item in payload.items)

    # E. Enregistrement en base de données avec statut PENDING
    po = PurchaseOrder(
        customer_id=customer.id,
        po_number=po_number,
        status=OrderStatus.PENDING,
        total_amount=total_amount
    )
    
    db.add(po)
    await db.commit()
    await db.refresh(po)

    return {
        "message": "Bon de commande créé avec succès !",
        "order_id": po.id,
        "po_number": po.po_number,
        "customer": customer.company_name,
        "total_amount": po.total_amount,
        "status": po.status
    }


@router.post("/upload-document", status_code=status.HTTP_201_CREATED)
async def upload_order_document(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db)
):
    """3. Traiter un document uploadé (PDF / Word) pour générer un PO de manière sécurisée"""
    try:
        contents = await file.read()
        extracted_text = extract_text_from_file(contents, file.filename)

        if not extracted_text or not extracted_text.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, 
                detail="Le fichier est vide ou le texte n'a pas pu être extrait."
            )

        # Extraction basique du client
        client_name = "Client Importé"
        lines = [line.strip() for line in extracted_text.split("\n") if line.strip()]
        
        for line in lines:
            if any(k in line.lower() for k in ["client", "société", "company", "marjane", "label'vie", "glace", "aswak"]):
                client_name = line.split(":")[-1].strip() if ":" in line else line
                break
        if client_name == "Client Importé" and len(lines) > 0:
            client_name = lines[0][:40]

        # Recherche des montants
        numbers = re.findall(r'\b\d+(?:[\.,]\d{1,2})?\b', extracted_text)
        floats = []
        for num in numbers:
            try:
                floats.append(float(num.replace(',', '.')))
            except ValueError:
                pass
        total_amount = max(floats) if floats else 1000.0

        # Recherche ou création automatique du client en respectant les contraintes Supabase
        result = await db.execute(
            select(Customer).filter(Customer.company_name.ilike(f"%{client_name}%"))
        )
        customer = result.scalars().first()

        if not customer:
            clean_code = "".join(e for e in client_name if e.isalnum()).upper()[:6] or "CUST"
            customer_code = f"C-{clean_code}-{str(uuid.uuid4())[:4].upper()}"

            detected_city = "Tétouan"
            for city_candidate in ["Rabat", "Tétouan", "Casablanca", "Tanger", "Martil", "Agadir"]:
                if city_candidate.lower() in client_name.lower() or city_candidate.lower() in extracted_text.lower():
                    detected_city = city_candidate.upper()
                    break

            customer = Customer(
                code=customer_code,
                company_name=client_name,
                email=f"contact@{clean_code.lower()}.com",
                address="Zone Industrielle",
                city=detected_city,
                ice="000000000000000"
            )
            db.add(customer)
            await db.commit()
            await db.refresh(customer)

        # Génération du numéro de PO
        raw_prefix = getattr(customer, "code", None) or customer.company_name[:3]
        prefix = "".join(e for e in raw_prefix if e.isalnum()).upper() or "PO"
        timestamp = datetime.now().strftime("%Y%m%d-%H%M")
        po_number = f"PO-{prefix}-{timestamp}"

        po = PurchaseOrder(
            customer_id=customer.id,
            po_number=po_number,
            status=OrderStatus.PENDING,
            total_amount=total_amount
        )

        db.add(po)
        await db.commit()
        await db.refresh(po)

        return {
            "message": "Document importé avec succès !",
            "order_id": po.id,
            "po_number": po.po_number,
            "customer": customer.company_name,
            "total_amount": po.total_amount,
            "status": po.status
        }

    except HTTPException:
        raise
    except Exception as e:
        print(f"Erreur serveur lors de la lecture du fichier : {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Erreur lors du traitement du fichier : {str(e)}"
        )


@router.post("/{order_id}/convert-to-bl", status_code=status.HTTP_201_CREATED)
async def convert_po_to_bl(
    order_id: str, 
    driver_name: str | None = "Chauffeur Glacematic",
    db: AsyncSession = Depends(get_db)
):
    """4. Transformer un Bon de Commande (PO) en Bon de Livraison (BL) au format Glacematic"""
    
    # A. Récupérer le PO avec le client associé
    result = await db.execute(
        select(PurchaseOrder)
        .options(selectinload(PurchaseOrder.customer))
        .filter(PurchaseOrder.id == order_id)
    )
    po = result.scalars().first()

    if not po:
        raise HTTPException(status_code=404, detail="Bon de commande introuvable")

    # B. Générer le numéro de BL au format officiel (ex: "04/2026")
    current_year = datetime.now().strftime("%Y")
    count_result = await db.execute(select(DeliveryNote))
    bl_count = len(count_result.scalars().all()) + 1
    bl_number = f"{bl_count:02d}/{current_year}"

    # C. Créer l'enregistrement Bon de Livraison dans la base
    bl = DeliveryNote(
        order_id=po.id,
        bl_number=bl_number,
        driver_name=driver_name,
        status=DeliveryStatus.PREPARING if hasattr(DeliveryStatus, "PREPARING") else "EN_COURS"
    )
    
    # D. Mettre à jour le statut du PO
    po.status = OrderStatus.CONFIRMED if hasattr(OrderStatus, "CONFIRMED") else "CONFIRMED"

    db.add(bl)
    await db.commit()
    await db.refresh(bl)

    return {
        "message": "Bon de livraison généré avec succès !",
        "bl_id": bl.id,
        "bl_number": bl.bl_number,
        "po_number": po.po_number,
        "customer_name": po.customer.company_name if po.customer else "Client"
    }